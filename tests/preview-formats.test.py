import importlib.util
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image

HELPER = Path(__file__).resolve().parents[1] / 'shell/bingux/preview-document.py'
spec = importlib.util.spec_from_file_location('preview', HELPER)
preview = importlib.util.module_from_spec(spec)
spec.loader.exec_module(preview)


class ExtendedPreviewTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.path = Path(temporary.name)
        environment = patch.dict('os.environ', {'XDG_CACHE_HOME': str(self.path / 'cache')})
        environment.start()
        self.addCleanup(environment.stop)

    def test_office_pages_preserve_document_content(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import Workbook
        document = Document()
        document.add_heading('Document heading', 0)
        document.add_paragraph('Word document preview')
        document.save(self.path / 'word.docx')
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = 'Slide heading'
        slides.save(self.path / 'slides.pptx')
        workbook = Workbook()
        workbook.active.append(['Name', 'Value'])
        workbook.active.append(['Spreadsheet preview', 42])
        workbook.save(self.path / 'sheet.xlsx')
        for name in ['word.docx', 'slides.pptx', 'sheet.xlsx']:
            with self.subTest(name=name):
                path = self.path / name
                original = path.read_bytes()
                result = preview.metadata(path)
                self.assertEqual(result['kind'], 'pdf')
                self.assertGreater(len(result['pages']), 0)
                self.assertTrue(preview.render(path, 0, 320)['image'].startswith('data:image/png;base64,'))
                converted = preview.formats.office_pdf(path)
                self.assertEqual(converted, preview.formats.office_pdf(path))
                pdf = preview.Poppler.Document.new_from_file(converted.as_uri(), None)
                self.assertTrue(pdf.get_page(0).get_text().strip())
                self.assertEqual(path.read_bytes(), original)

    def test_slow_converter_is_stopped(self):
        with self.assertRaisesRegex(ValueError, 'too long'):
            preview.formats.run_tool([sys.executable, '-c', 'import time; time.sleep(10)'], timeout=0.02)

    def test_database_pagination_quotes_and_read_only(self):
        path = self.path / 'example.sqlite'
        connection = sqlite3.connect(path)
        connection.execute('CREATE TABLE "odd""name" (id INTEGER, value TEXT, data BLOB)')
        connection.executemany('INSERT INTO "odd""name" VALUES (?, ?, ?)', [(index, 'Row ' + str(index), b'\x00\x01') for index in range(125)])
        connection.commit()
        connection.close()
        original = path.read_bytes()
        result = preview.metadata(path)
        self.assertEqual(result['kind'], 'database')
        self.assertEqual(len(result['rows']), 100)
        self.assertTrue(result['hasMore'])
        self.assertEqual(result['rows'][0][2], '[2 bytes]')
        second = preview.formats.database(path, 'odd"name', 100)
        self.assertEqual(len(second['rows']), 25)
        self.assertEqual(second['rows'][0][0], '100')
        self.assertFalse(second['hasMore'])
        self.assertEqual(path.read_bytes(), original)

    def test_formatted_markdown_and_html_do_not_execute_or_fetch(self):
        for name, text in [('note.md', '# Heading\n\n**Bold** text\n\n![external](https://example.com/image.png)'), ('page.html', '<h1>Heading</h1><script>alert(1)</script><img src="https://example.com/image.png"><p>Content</p>')]:
            path = self.path / name
            path.write_text(text)
            result = preview.metadata(path)
            self.assertIn('<h1>Heading</h1>', result['renderedText'])
            self.assertNotIn('https://', result['renderedText'])
            self.assertNotIn('alert(1)', result['renderedText'])

    def test_gif_frames_and_video_metadata(self):
        path = self.path / 'animated.gif'
        frames = [Image.new('RGB', (64, 32), colour) for colour in ['red', 'blue']]
        frames[0].save(path, save_all=True, append_images=frames[1:], duration=100, loop=0)
        self.assertEqual(preview.metadata(path)['kind'], 'animation')
        path = self.path / 'clip.mp4'
        subprocess.run(['ffmpeg', '-v', 'error', '-f', 'lavfi', '-i', 'color=c=red:s=160x90:d=1', '-c:v', 'mpeg4', str(path)], check=True)
        result = preview.metadata(path)
        self.assertEqual(result['kind'], 'video')
        self.assertEqual((result['width'], result['height']), (160, 90))
        self.assertGreater(result['duration'], 0)
        self.assertTrue(preview.formats.video_poster(path).startswith('data:image/png;base64,'))

    def test_size_limit_applies_to_media_database_and_direct_table_reads(self):
        for suffix in ['.mp4', '.sqlite', '.gif', '.docx']:
            path = self.path / ('large' + suffix)
            with path.open('wb') as stream:
                stream.truncate(preview.MAX_FILE_BYTES + 1)
            with self.assertRaisesRegex(ValueError, '20 MB limit'):
                preview.metadata(path)
        with self.assertRaisesRegex(ValueError, '20 MB limit'):
            preview.formats.database(self.path / 'large.sqlite')


if __name__ == '__main__':
    unittest.main()
